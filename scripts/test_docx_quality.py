import docx

doc = docx.Document("VeLiS_RAG_Architecture_and_Implementation_Plan_Revised.docx")
full_text = []
for p in doc.paragraphs:
    full_text.append(p.text)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            full_text.append(cell.text)
corpus = "\n".join(full_text)

placeholders = [
    "[Author’s project motivation]",
    "[Local hardware configuration]",
    "[Supervisor feedback]",
    "[Observed Phase 1 results]",
    "[Author’s implementation decision and reason]",
]

print("=== Placeholder Check ===")
for ph in placeholders:
    straight = ph.replace("’", "'")
    c1 = corpus.count(ph)
    c2 = corpus.count(straight)
    print(f"{ph}: {c1 + c2} occurrences")

words = ["strict", "critical", "guarantees", "scientific rigor", "deterministic", "bulletproof", "catastrophic"]
print("\n=== Term Frequency Check ===")
for w in words:
    print(f"{w}: {corpus.lower().count(w)}")
